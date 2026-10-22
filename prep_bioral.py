import openpyxl, json, re, sys, unicodedata, os
from collections import defaultdict

SRC = sys.argv[1]
SC   = '/tmp/claude-0/-home-claude/0181ac0f-dbc6-5800-b9bf-408f428f8d00/scratchpad/'

def sa(s):  # sem acento, maiusculo
    return ''.join(c for c in unicodedata.normalize('NFD', str(s or '')) if unicodedata.category(c) != 'Mn').upper().strip()

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# ---- especialidades por estabelecimento, vindas da aba de profissionais ----
esp_por = defaultdict(set)
prof_por = defaultdict(int)
for r in wb['Profissionais'].iter_rows(min_row=2, values_only=True):
    if not r or not r[0]:
        continue
    k = (sa(r[0]), sa(r[9]), str(r[11] or '').strip())      # estabelecimento, endereco, cep
    prof_por[k] += 1
    if r[7]:
        for e in str(r[7]).split('|'):
            e = e.strip()
            if e:
                esp_por[k].add(e)

esp_set = sorted({e for v in esp_por.values() for e in v})
esp_ix  = {e: i for i, e in enumerate(esp_set)}

TIPO = {'PESSOA JURIDICA': 0, 'PESSOA FISICA': 1}

def fone(fixo, cel, zap):
    for v in (fixo, cel, zap):
        if v:
            return str(v).split(',')[0].strip()
    return ''

pontos, consultas = [], []
for r in wb['Estabelecimentos'].iter_rows(min_row=2, values_only=True):
    if not r or not r[0]:
        continue
    nome   = (r[3] or r[0] or '').strip()          # nome fantasia, senao estabelecimento
    end    = str(r[5] or '').strip()
    bairro = str(r[6] or '').strip()
    cep    = str(r[7] or '').strip()
    k = (sa(r[0]), sa(end), cep)
    p = {
        'n': nome,
        'e': end,
        'b': bairro,
        'c': cep,
        'f': fone(r[11], r[12], r[13]),
        't': TIPO.get(sa(r[1]), 0),
        'np': int(r[14]) if (r[14] and str(r[14]).isdigit()) else prof_por.get(k, 0),
        'es': sorted(esp_ix[e] for e in esp_por.get(k, ())),
    }
    pontos.append(p)
    consultas.append(end + (', ' + bairro if bairro else '') + ', ' + cep + ', Sao Paulo')

os.makedirs('_geo', exist_ok=True)
json.dump(consultas, open('_geo/bioral_q.json', 'w'), ensure_ascii=False)
json.dump({'pontos': pontos, 'esp': esp_set}, open(SC + 'bioral_prep.json', 'w'), ensure_ascii=False)

sem_esp = sum(1 for p in pontos if not p['es'])
print('estabelecimentos :', len(pontos))
print('  pessoa juridica:', sum(1 for p in pontos if p['t'] == 0))
print('  pessoa fisica  :', sum(1 for p in pontos if p['t'] == 1))
print('profissionais    :', sum(p['np'] for p in pontos))
print('sem especialidade:', sem_esp)
print('sem telefone     :', sum(1 for p in pontos if not p['f']))
print('especialidades   :', len(esp_set))
print('consultas geo    :', len(consultas), '| bytes:', os.path.getsize('_geo/bioral_q.json'))
print('exemplo          :', consultas[0])
